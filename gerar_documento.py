# -*- coding: utf-8 -*-
"""Gera o documento do projeto (.docx): o que foi feito, como funciona o
score, regras usadas e explicacao de seguranca."""

import os

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DIR_BASE = os.path.dirname(os.path.abspath(__file__))
CAMINHO_DOCX = os.path.join(DIR_BASE, "documento_projeto.docx")

AZUL = RGBColor(0x1F, 0x3A, 0x5F)


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def tabela(doc, cabecalho, linhas, larguras=None):
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, texto in enumerate(cabecalho):
        celula = t.rows[0].cells[i]
        celula.text = ""
        p = celula.paragraphs[0]
        r = p.add_run(texto)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)
        set_cell_shading(celula, "1F3A5F")
    for linha in linhas:
        celulas = t.add_row().cells
        for i, valor in enumerate(linha):
            celulas[i].text = ""
            p = celulas[i].paragraphs[0]
            r = p.add_run(str(valor))
            r.font.size = Pt(9.5)
    if larguras:
        for row in t.rows:
            for i, w in enumerate(larguras):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def montar():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    titulo = doc.add_heading("Sistema de Risco IoT", level=0)
    for run in titulo.runs:
        run.font.color.rgb = AZUL
    sub = doc.add_paragraph("Documento do projeto")
    sub.runs[0].italic = True
    doc.add_paragraph("Disciplina: Cibersegurança Cognitiva - Challenge Sprint 3")
    doc.add_paragraph("Professor: Gilberto Vieira Branco")
    doc.add_paragraph("Projeto: Sompo Seguros")

    doc.add_heading("1. O que fizemos", level=1)
    doc.add_paragraph(
        "Montamos um sistema completo que integra as etapas pedidas no enunciado: geração "
        "de dados IoT, envio para um servidor, armazenamento, análise de risco e um "
        "resultado organizado no final. O fluxo segue a ordem pedida: IoT, Servidor, Dados, "
        "Análise, Resultado."
    )
    doc.add_paragraph(
        "O enunciado pede pra reaproveitar o código das Sprints 1 e 2 da disciplina. A gente "
        "não tinha esse código em mãos na hora de montar esse projeto, então construímos do "
        "zero cobrindo o mesmo fluxo que as Sprints 1 e 2 provavelmente já cobriam separado "
        "(gerar e enviar dados IoT na Sprint 1, analisar e gerar alertas na Sprint 2), agora "
        "tudo integrado num sistema só."
    )
    doc.add_paragraph(
        "O sistema simula sensores de temperatura e umidade em 5 equipamentos rurais. A cada "
        "execução, o gerador cria leituras normais e mistura de propósito algumas leituras "
        "problemáticas (temperatura alta, umidade baixa, dado incompleto e uma repetição), "
        "pra garantir que dá pra ver todas as regras de risco funcionando."
    )

    doc.add_heading("2. Como o sistema está organizado", level=1)
    tabela(doc, ["Etapa", "Onde está no código", "O que faz"], [
        ["IoT", "iot/gerador.py", "Gera as leituras simuladas dos sensores"],
        ["Envio", "iot/cliente.py", "Envia cada leitura para o servidor via HTTP POST"],
        ["Servidor", "servidor/app.py", "Recebe, valida e armazena as leituras (arquivo local)"],
        ["Análise", "analise/regras.py, motor.py", "Aplica as regras e calcula o score por equipamento"],
        ["Segurança", "analise/seguranca.py", "Identifica dados inválidos e leituras suspeitas"],
        ["Resultado", "saida/relatorio.py", "Mostra o resumo e exporta em JSON/CSV"],
    ], larguras=[3, 5, 8])

    doc.add_heading("3. Como funciona o score de risco", level=1)
    doc.add_paragraph(
        "Cada leitura que chega passa por um conjunto de regras. Toda regra que dispara "
        "soma pontos ao score do equipamento correspondente. O score final de um "
        "equipamento é a soma dos pontos de todas as leituras dele na rodada."
    )
    tabela(doc, ["Regra", "Condição", "Pontos"], [
        ["Temperatura alta", "temperatura acima de 35°C", "+2"],
        ["Umidade baixa", "umidade abaixo de 30%", "+2"],
        ["Dados incompletos", "campo obrigatório ausente na leitura", "+3"],
        ["Repetição", "mesma temperatura e umidade em leituras seguidas do mesmo equipamento", "+2"],
    ], larguras=[3.5, 9, 2.5])
    doc.add_paragraph("A classificação final segue a tabela que o professor deu no enunciado:")
    tabela(doc, ["Score", "Nível"], [
        ["0 a 2", "BAIXO"],
        ["3 a 5", "MEDIO"],
        ["6 ou mais", "ALTO"],
    ], larguras=[8, 8])

    doc.add_heading("4. Segurança", level=1)

    doc.add_heading("4.1 Como evitamos dados inválidos", level=2)
    doc.add_paragraph(
        "A validação acontece no próprio servidor, antes de qualquer dado ser guardado "
        "(função validar_payload em servidor/app.py). O servidor rejeita, com resposta HTTP "
        "400, qualquer leitura que: não seja um objeto JSON válido; esteja sem os campos "
        "obrigatórios (equipamento, temperatura, timestamp); ou tenha temperatura ou "
        "umidade fora de uma faixa fisicamente plausível (temperatura entre -30°C e 80°C, "
        "umidade entre 0% e 100%). Assim, o que chega na análise já passou por um primeiro "
        "filtro de sanidade."
    )

    doc.add_heading("4.2 Como identificamos dados suspeitos", level=2)
    doc.add_paragraph(
        "Nem todo dado suspeito é claramente inválido, então a camada de análise "
        "(analise/seguranca.py) faz umas checagens a mais depois que os dados já foram "
        "guardados: confere de novo os limites físicos dos valores (uma segunda camada de "
        "proteção, caso algo escape da validação do servidor); sinaliza leituras com "
        "timestamp no futuro; e sinaliza leituras com um atraso muito grande entre o "
        "horário que o sensor registrou e o horário que o servidor recebeu, o que pode "
        "indicar reenvio de um dado antigo (replay) ou relógio do sensor adulterado. A "
        "própria regra de repetição do score de risco também serve como indício de dado "
        "suspeito: um sensor que manda exatamente a mesma leitura duas vezes seguidas pode "
        "estar com defeito ou sendo usado pra inflar os dados de propósito."
    )

    doc.add_heading("4.3 Como melhorar a segurança em produção", level=2)
    doc.add_paragraph(
        "Esse é um MVP acadêmico e roda local sem essas proteções, mas num sistema real de "
        "produção a gente precisaria de pelo menos:"
    )
    for item in [
        "Comunicação via HTTPS/TLS entre os dispositivos IoT e o servidor, pra impedir que "
        "os dados sejam lidos ou alterados no meio do caminho.",
        "Autenticação por dispositivo (API key ou certificado por sensor), pra o servidor "
        "só aceitar dados de equipamentos cadastrados e dar pra revogar o acesso de um "
        "dispositivo comprometido.",
        "Limite de taxa de envio (rate limiting), pra impedir que um sensor com defeito ou "
        "malicioso sobrecarregue o servidor ou infle os dados.",
        "Assinatura ou hash de cada leitura, pra detectar se o conteúdo foi alterado entre "
        "o sensor e o servidor.",
        "Log de auditoria de tudo que é rejeitado, pra acompanhar tentativas repetidas de "
        "envio de dados inválidos, o que pode indicar um sensor com defeito ou um ataque.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("5. Requisitos técnicos atendidos", level=1)
    tabela(doc, ["Requisito", "Onde está"], [
        ["Listas", "gerador.py (lista de leituras), motor.py (agrupamento por equipamento)"],
        ["Estruturas de repetição", "loops em motor.py, relatorio.py e cliente.py"],
        ["Funções", "todo o fluxo é dividido em funções por responsabilidade"],
        ["Sistema integrado (IoT -> Servidor -> Dados -> Análise -> Resultado)", "main.py orquestra as 5 etapas numa única execução"],
        ["Saída organizada (não só print)", "JSON e CSV exportados em output/"],
    ], larguras=[9, 7])

    doc.save(CAMINHO_DOCX)
    print("Documento gerado em:", CAMINHO_DOCX)


if __name__ == "__main__":
    montar()
